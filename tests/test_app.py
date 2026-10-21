import io
import unittest
from types import SimpleNamespace
from unittest.mock import patch

import pymupdf
from docx import Document
from PIL import Image

import app as server
from services import MAX_TEXT, MAX_UPLOAD, detect_language, extract_file


def completion(text, reason="stop"):
    return SimpleNamespace(
        choices=[SimpleNamespace(message=SimpleNamespace(content=text), finish_reason=reason)]
    )


def docx_bytes():
    document = Document()
    document.add_paragraph("First paragraph")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Left", "Right"
    document.add_paragraph("Last paragraph")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def pdf_bytes(pages=1, scanned=False):
    with pymupdf.open() as document:
        for _ in range(pages):
            page = document.new_page()
            if scanned:
                image = Image.new("RGB", (200, 100), "white")
                buffer = io.BytesIO()
                image.save(buffer, format="PNG")
                page.insert_image(pymupdf.Rect(0, 0, 200, 100), stream=buffer.getvalue())
            else:
                page.insert_text((72, 72), "Hello from PDF")
        return document.tobytes()


class AppTests(unittest.TestCase):
    def setUp(self):
        self.client = server.app.test_client()

    def test_translation_validates_body_and_settings(self):
        for body in [
            None,
            [],
            "text",
            {"text": 123},
            {"text": ""},
            {"text": "x", "source": []},
            {"text": "x", "target": {}},
            {"text": "x", "tone": "invalid"},
            {"text": "x" * (MAX_TEXT + 1)},
        ]:
            with self.subTest(body=str(body)[:50]):
                result = self.client.post("/api/translate", json=body)
                self.assertEqual(result.status_code, 400)
                self.assertIn("error", result.json)

    def test_same_language_and_long_translation(self):
        result = self.client.post(
            "/api/translate", json={"text": "สวัสดี", "source": "th", "target": "th"}
        )
        self.assertEqual(result.json["translation"], "สวัสดี")
        text = ("First paragraph.\n" * 400) + "The last line."
        self.assertEqual("".join(server.chunks(text)), text)
        with patch.object(server, "get_client") as client:
            client.return_value.chat.completions.create.side_effect = lambda **kw: completion(
                kw["messages"][1]["content"].strip()
            )
            result = self.client.post(
                "/api/translate", json={"text": text, "source": "en", "target": "th"}
            )
            self.assertEqual(result.status_code, 200)
            self.assertIn("The last line.", result.json["translation"])
            self.assertGreater(client.return_value.chat.completions.create.call_count, 1)

    def test_truncated_response_is_error(self):
        with patch.object(server, "get_client") as client:
            client.return_value.chat.completions.create.return_value = completion(
                "Partial", "length"
            )
            result = self.client.post("/api/translate", json={"text": "hello", "target": "th"})
            self.assertEqual(result.status_code, 503)
            self.assertNotIn("translation", result.json)

    def test_document_order_and_formats(self):
        self.assertEqual(
            extract_file(docx_bytes(), "test.docx", "document", None),
            "First paragraph\nLeft\tRight\nLast paragraph",
        )
        self.assertEqual(extract_file(pdf_bytes(), "test.pdf", "document", None), "Hello from PDF")
        self.assertEqual(extract_file("\ufeffสวัสดี".encode(), "test.txt", "document", None), "สวัสดี")
        for extension in ["txt", "md", "csv"]:
            with self.subTest(extension=extension):
                result = self.client.post(
                    "/api/extract",
                    data={"kind": "document", "file": (io.BytesIO(b"hello"), f"test.{extension}")},
                )
                self.assertEqual(result.json["text"], "hello")

    def test_scanned_pdf_calls_ocr(self):
        with patch("services.read_image", return_value="Scanned text") as ocr:
            self.assertEqual(
                extract_file(pdf_bytes(scanned=True), "scan.pdf", "document", None), "Scanned text"
            )
            ocr.assert_called_once()

    def test_bad_files_are_json_errors(self):
        for kind, name, data in [
            ("image", "bad.png", b"not an image"),
            ("document", "bad.pdf", b"bad"),
            ("document", "bad.docx", b"bad"),
            ("document", "bad.txt", b"\xff\xff"),
            ("document", "binary.txt", b"a\x00b"),
            ("document", "test.exe", b"test"),
            ("document", "long.txt", b"x" * (MAX_TEXT + 1)),
            ("unknown", "x.txt", b"x"),
        ]:
            with self.subTest(name=name):
                result = self.client.post(
                    "/api/extract", data={"kind": kind, "file": (io.BytesIO(data), name)}
                )
                self.assertEqual(result.status_code, 400)
                self.assertIn("error", result.json)
        result = self.client.post(
            "/api/extract", data={"file": (io.BytesIO(b"x" * (MAX_UPLOAD + 1)), "huge.txt")}
        )
        self.assertEqual(result.status_code, 413)
        result.request.input_stream.close()
        result.close()
        result = self.client.post(
            "/api/extract", data={"file": (io.BytesIO(pdf_bytes(pages=21)), "long.pdf")}
        )
        self.assertEqual(result.status_code, 400)
        result = self.client.post("/api/extract", data={})
        self.assertEqual(result.status_code, 400)

    def test_speech_and_transcription(self):
        with patch.object(server, "create_speech") as speech:
            speech.return_value = b"ID3test-audio"
            result = self.client.post("/api/speech", json={"text": "สวัสดี", "language": "auto"})
            self.assertEqual(result.status_code, 200)
            self.assertEqual(result.content_type, "audio/mpeg")
            speech.assert_called_once_with("สวัสดี", "th")
        with patch.object(server, "get_client") as client:
            client.return_value.audio.transcriptions.create.return_value = SimpleNamespace(
                text="สวัสดี"
            )
            result = self.client.post(
                "/api/transcribe",
                data={"language": "th", "file": (io.BytesIO(b"audio"), "recording.webm")},
            )
            self.assertEqual(result.json, {"text": "สวัสดี"})
            self.assertEqual(
                client.return_value.audio.transcriptions.create.call_args.kwargs["language"], "th"
            )
        self.assertEqual(
            self.client.post("/api/speech", json={"text": "hi", "language": []}).status_code, 400
        )
        self.assertEqual(detect_language("안녕하세요"), "ko")
        self.assertEqual(detect_language("こんにちは"), "ja")


if __name__ == "__main__":
    unittest.main()
